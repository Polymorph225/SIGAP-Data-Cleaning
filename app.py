import streamlit as st
import pandas as pd
import re
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
#  PAGE CONFIG
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Data Cleaning SIGAP Bojonegoro",
    page_icon="🧹",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────
#  CUSTOM CSS
# ─────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap');

html, body, [class*="css"] {
    font-family: 'Plus Jakarta Sans', sans-serif;
}

/* ── Background ── */
.stApp {
    background: linear-gradient(135deg, #0f1623 0%, #1a2744 50%, #0f1623 100%);
    min-height: 100vh;
}

/* ── Sidebar ── */
section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #111827 0%, #1e2d4a 100%);
    border-right: 1px solid rgba(99,179,237,0.15);
}
section[data-testid="stSidebar"] * { color: #e2e8f0 !important; }
section[data-testid="stSidebar"] .stSelectbox label,
section[data-testid="stSidebar"] .stCheckbox label { color: #94a3b8 !important; }

/* ── Header ── */
.main-header {
    background: linear-gradient(135deg, rgba(30,58,138,0.6) 0%, rgba(14,165,233,0.15) 100%);
    border: 1px solid rgba(99,179,237,0.3);
    border-radius: 16px;
    padding: 2rem 2.5rem;
    margin-bottom: 1.5rem;
    position: relative;
    overflow: hidden;
}
.main-header::before {
    content: '';
    position: absolute;
    top: -50%;
    right: -10%;
    width: 300px;
    height: 300px;
    background: radial-gradient(circle, rgba(14,165,233,0.12) 0%, transparent 70%);
    border-radius: 50%;
}
.main-header h1 {
    color: #f0f9ff;
    font-size: 2.2rem;
    font-weight: 800;
    margin: 0 0 0.25rem 0;
    letter-spacing: -0.03em;
}
.main-header p {
    color: #7dd3fc;
    font-size: 0.95rem;
    margin: 0;
    font-weight: 500;
}

/* ── Metric cards ── */
.metric-row {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 1rem;
    margin-bottom: 1.5rem;
}
.metric-card {
    background: rgba(30,41,59,0.8);
    border: 1px solid rgba(99,179,237,0.2);
    border-radius: 12px;
    padding: 1.25rem 1.5rem;
    text-align: center;
    backdrop-filter: blur(8px);
}
.metric-card .label {
    color: #64748b;
    font-size: 0.75rem;
    font-weight: 600;
    text-transform: uppercase;
    letter-spacing: 0.08em;
    margin-bottom: 0.5rem;
}
.metric-card .value {
    color: #e2e8f0;
    font-size: 1.9rem;
    font-weight: 800;
    line-height: 1;
    font-family: 'JetBrains Mono', monospace;
}
.metric-card .value.good { color: #34d399; }
.metric-card .value.warn { color: #fbbf24; }
.metric-card .value.bad  { color: #f87171; }

/* ── Section card ── */
.section-card {
    background: rgba(15,22,35,0.7);
    border: 1px solid rgba(99,179,237,0.15);
    border-radius: 14px;
    padding: 1.5rem;
    margin-bottom: 1rem;
    backdrop-filter: blur(6px);
}
.section-title {
    color: #7dd3fc;
    font-size: 0.8rem;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 1rem;
    display: flex;
    align-items: center;
    gap: 0.5rem;
}

/* ── Issue badges ── */
.badge {
    display: inline-block;
    padding: 0.2rem 0.65rem;
    border-radius: 999px;
    font-size: 0.72rem;
    font-weight: 600;
    font-family: 'JetBrains Mono', monospace;
}
.badge-red   { background: rgba(248,113,113,0.15); color: #f87171; border: 1px solid rgba(248,113,113,0.3); }
.badge-yellow{ background: rgba(251,191,36,0.12);  color: #fbbf24; border: 1px solid rgba(251,191,36,0.3); }
.badge-green { background: rgba(52,211,153,0.12);  color: #34d399; border: 1px solid rgba(52,211,153,0.3); }
.badge-blue  { background: rgba(99,179,237,0.12);  color: #63b3ed; border: 1px solid rgba(99,179,237,0.3); }

/* ── Log item ── */
.log-item {
    display: flex;
    align-items: flex-start;
    gap: 0.6rem;
    padding: 0.6rem 0;
    border-bottom: 1px solid rgba(99,179,237,0.07);
    font-size: 0.85rem;
    color: #94a3b8;
}
.log-icon { font-size: 1rem; flex-shrink: 0; margin-top: 1px; }

/* ── Upload zone ── */
.upload-hint {
    background: rgba(14,165,233,0.06);
    border: 2px dashed rgba(14,165,233,0.25);
    border-radius: 12px;
    padding: 1.5rem;
    text-align: center;
    color: #475569;
    font-size: 0.9rem;
    margin-bottom: 1rem;
}

/* ── Dataframe ── */
[data-testid="stDataFrame"] { border-radius: 10px; overflow: hidden; }

/* ── Buttons ── */
.stButton > button {
    background: linear-gradient(135deg, #1d4ed8, #0ea5e9) !important;
    color: white !important;
    border: none !important;
    border-radius: 10px !important;
    font-weight: 700 !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-size: 0.9rem !important;
    padding: 0.6rem 1.5rem !important;
    transition: all 0.2s !important;
    box-shadow: 0 4px 15px rgba(14,165,233,0.25) !important;
}
.stButton > button:hover {
    transform: translateY(-1px) !important;
    box-shadow: 0 6px 20px rgba(14,165,233,0.4) !important;
}

/* ── Download button ── */
.stDownloadButton > button {
    background: linear-gradient(135deg, #065f46, #10b981) !important;
    color: white !important;
    border: none !important;
    border-radius: 10px !important;
    font-weight: 700 !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    width: 100% !important;
    padding: 0.75rem !important;
    font-size: 1rem !important;
    box-shadow: 0 4px 15px rgba(16,185,129,0.3) !important;
}

/* ── Checkboxes & selects ── */
.stCheckbox > label { color: #cbd5e1 !important; font-size: 0.9rem !important; }
.stSelectbox > label { color: #94a3b8 !important; font-size: 0.8rem !important; font-weight: 600 !important; }
.stMultiSelect > label { color: #94a3b8 !important; font-size: 0.8rem !important; font-weight: 600 !important; }

/* ── Divider ── */
hr { border-color: rgba(99,179,237,0.1) !important; }

/* ── Expander ── */
.streamlit-expanderHeader {
    background: rgba(30,41,59,0.5) !important;
    border-radius: 8px !important;
    color: #7dd3fc !important;
}
</style>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────

def detect_issues(df: pd.DataFrame) -> dict:
    """Detect data quality issues and return a report dict."""
    issues = {}

    # Trailing comma/space patterns
    trailing_comma_cols = []
    for col in df.select_dtypes(include="object").columns:
        if df[col].dropna().astype(str).str.contains(r'\s*,\s*$', regex=True).any():
            trailing_comma_cols.append(col)
    if trailing_comma_cols:
        issues["trailing_comma"] = trailing_comma_cols

    # Duplicates
    dup_count = df.duplicated().sum()
    if dup_count:
        issues["duplicates"] = int(dup_count)

    # Missing values
    nulls = df.isnull().sum()
    null_cols = nulls[nulls > 0].to_dict()
    if null_cols:
        issues["nulls"] = null_cols

    # Placeholder values
    placeholder_cols = []
    for col in df.select_dtypes(include="object").columns:
        if df[col].dropna().astype(str).str.fullmatch(r'[-=]+').any():
            placeholder_cols.append(col)
    if placeholder_cols:
        issues["placeholders"] = placeholder_cols

    # Date columns
    date_cols = []
    for col in df.columns:
        if "tgl" in col.lower() or "tanggal" in col.lower() or "date" in col.lower() or "lahir" in col.lower():
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                date_cols.append(col)
    if date_cols:
        issues["date_cols"] = date_cols

    return issues


def clean_dataframe(df: pd.DataFrame, options: dict) -> tuple[pd.DataFrame, list]:
    """Apply cleaning steps based on options. Returns (cleaned_df, log)."""
    df = df.copy()
    log = []

    # 1. Remove duplicates
    if options.get("remove_duplicates"):
        before = len(df)
        df.drop_duplicates(inplace=True)
        removed = before - len(df)
        if removed:
            log.append(("🗑️", f"Hapus {removed} baris duplikat"))
        else:
            log.append(("✅", "Tidak ada baris duplikat"))

    # 2. Strip trailing commas/spaces from all object cols
    if options.get("strip_trailing_comma"):
        affected = []
        for col in df.select_dtypes(include="object").columns:
            mask = df[col].dropna().astype(str).str.contains(r'\s*,\s*$', regex=True)
            if mask.any():
                df[col] = df[col].astype(str).str.replace(r'\s*,\s*$', '', regex=True).str.strip()
                affected.append(col)
        if affected:
            log.append(("✂️", f"Hapus trailing koma pada: {', '.join(affected)}"))

    # 3. Strip extra internal spaces
    if options.get("strip_spaces"):
        for col in options.get("strip_spaces_cols", []):
            if col in df.columns:
                df[col] = df[col].astype(str).str.replace(r'\s+', '', regex=True)
        log.append(("🔤", "Hapus spasi berlebih di dalam NIK / No Penjamin"))

    # 4. Format dates
    if options.get("format_dates"):
        fmt = options.get("date_format", "%d/%m/%Y")
        for col in df.columns:
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                df[col] = pd.to_datetime(df[col]).dt.strftime(fmt)
                log.append(("📅", f"Format tanggal kolom '{col}' → {fmt}"))

    # 5. Standardize placeholder Keterangan
    if options.get("standardize_keterangan") and "Keterangan" in df.columns:
        placeholder_val = options.get("placeholder_replacement", "Tidak Ada Keterangan")
        df["Keterangan"] = df["Keterangan"].astype(str).str.strip()
        df["Keterangan"] = df["Keterangan"].replace(["-", "=", "nan", ""], placeholder_val)
        df["Keterangan"] = df["Keterangan"].apply(
            lambda x: placeholder_val if re.fullmatch(r'[-=\s]+', str(x)) else x
        )
        if options.get("uppercase_keterangan"):
            df["Keterangan"] = df["Keterangan"].str.upper()
        log.append(("📝", f"Standarisasi Keterangan: '-', '=' → '{placeholder_val}'"))

    # 6. Uppercase Nama
    if options.get("uppercase_nama") and "Nama" in df.columns:
        df["Nama"] = df["Nama"].str.upper().str.strip()
        log.append(("🔠", "Nama diubah ke UPPERCASE"))

    # 7. Title case Desa
    if options.get("titlecase_desa") and "Desa" in df.columns:
        df["Desa"] = df["Desa"].str.strip().str.title()
        log.append(("🏘️", "Desa diubah ke Title Case"))

    # 8. Fill nulls
    fill_map = options.get("fill_nulls", {})
    for col, fill_val in fill_map.items():
        if col in df.columns and fill_val:
            n = int(df[col].isnull().sum())
            if n:
                df[col] = df[col].fillna(fill_val)
                log.append(("📋", f"Isi {n} nilai kosong di '{col}' dengan '{fill_val}'"))

    # 9. Rename columns
    rename_map = options.get("rename_cols", {})
    if rename_map:
        df.rename(columns=rename_map, inplace=True)
        for old, new in rename_map.items():
            log.append(("✏️", f"Rename kolom '{old}' → '{new}'"))

    # 10. Drop selected columns
    drop_cols = [c for c in options.get("drop_cols", []) if c in df.columns]
    if drop_cols:
        df.drop(columns=drop_cols, inplace=True)
        log.append(("❌", f"Hapus kolom: {', '.join(drop_cols)}"))

    # 11. Sort
    sort_col = options.get("sort_by")
    if sort_col and sort_col in df.columns:
        df.sort_values(sort_col, inplace=True)
        df.reset_index(drop=True, inplace=True)
        log.append(("🔃", f"Urutkan berdasarkan '{sort_col}'"))

    return df, log


def to_excel_bytes(df: pd.DataFrame, sheet_name: str = "Data") -> bytes:
    """Export cleaned DataFrame to a nicely formatted Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    h_fill  = PatternFill('solid', start_color='1F4E79')
    h_font  = Font(bold=True, color='FFFFFF', name='Calibri', size=10)
    h_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    alt_fill = PatternFill('solid', start_color='EBF3FB')
    d_font   = Font(name='Calibri', size=9)
    d_align  = Alignment(vertical='center')
    thin = Side(style='thin', color='B0C4DE')
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header row
    for ci, col in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=ci, value=col)
        c.font = h_font; c.fill = h_fill
        c.alignment = h_align; c.border = bdr

    # Data rows
    df_str = df.fillna('-')
    for ri, row in enumerate(df_str.itertuples(index=False), 2):
        fill = alt_fill if ri % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = d_font; c.fill = fill
            c.alignment = d_align; c.border = bdr

    # Auto column widths (capped)
    for ci, col in enumerate(df.columns, 1):
        max_len = max(
            len(str(col)),
            df_str.iloc[:, ci-1].astype(str).str.len().max() if len(df_str) else 0
        )
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 45)

    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 32

    # Summary sheet
    ws2 = wb.create_sheet('Info')
    summary = [
        ['File diekspor oleh RME Data Cleaner'],
        ['Total baris:', len(df)],
        ['Total kolom:', len(df.columns)],
    ]
    for ri, row in enumerate(summary, 1):
        for ci, val in enumerate(row, 1):
            ws2.cell(row=ri, column=ci, value=val)
    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



# ─────────────────────────────────────────────
#  PEMBACAAN & PENGGABUNGAN BERKAS
# ─────────────────────────────────────────────

# Pemetaan kolom ekspor RME Sintesa Emas ke kolom masukan SIGAP-Bojonegoro.
KOLOM_SIGAP = {
    "Tgl":        "tanggal_kunjungan",
    "No RM":      "no_rm",
    "Usia":       "umur",
    "L/LP":       "jenis_kelamin",
    "Unit":       "poli",
    "Diagnosis":  "diagnosa",
    "Cara Bayar": "pembiayaan",
    "Desa":       "desa",
}


def _baca_teks_berpemisah(data: bytes) -> pd.DataFrame:
    """Ekspor RME kerap berekstensi .xls padahal isinya teks berpemisah tab.

    Fungsi ini menebak pemisahnya dari baris judul, lalu membaca seluruh kolom
    sebagai teks supaya nomor rekam medis dan NIK tidak berubah jadi angka.
    """
    teks = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            teks = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if teks is None:
        raise ValueError("Penyandian karakter berkas tidak dikenali.")

    baris_judul = teks.split("\n", 1)[0]
    pemisah = max(("\t", ";", "|", ","), key=baris_judul.count)
    if baris_judul.count(pemisah) == 0:
        raise ValueError("Tidak ditemukan pemisah kolom pada baris judul.")

    return pd.read_csv(
        io.StringIO(teks), sep=pemisah, dtype=str,
        quotechar='"', engine="python", keep_default_na=True,
    )


def baca_berkas(nama: str, data: bytes) -> pd.DataFrame:
    """Baca satu berkas apa pun bentuknya: csv, xlsx, xls asli, atau xls palsu."""
    n = nama.lower()
    if n.endswith(".csv"):
        return _baca_teks_berpemisah(data)
    if n.endswith(".xlsx"):
        return pd.read_excel(io.BytesIO(data), dtype=str, engine="openpyxl")
    # .xls — bisa Excel lama yang sebenarnya, bisa pula teks berpemisah
    try:
        return pd.read_excel(io.BytesIO(data), dtype=str, engine="xlrd")
    except Exception:
        return _baca_teks_berpemisah(data)


def _ke_tanggal(seri: pd.Series) -> pd.Series:
    """Ubah kolom tanggal menjadi datetime tanpa salah tebak hari/bulan.

    `dayfirst=True` berbahaya untuk tanggal ISO: pandas menebak formatnya dari
    nilai pertama, sehingga "2026-04-01" bisa terbaca 1 April atau 4 Januari,
    dan sisanya yang tanggalnya di atas 12 gagal diurai. Di sini setiap format
    yang mungkin dicoba satu per satu, lalu yang paling banyak berhasil dipakai.
    """
    teks = seri.astype(str).str.strip()
    formats = ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y",
               "%Y/%m/%d", "%d/%m/%y", "%m/%d/%Y"]
    terbaik, jml_terbaik = None, -1
    for fmt in formats:
        t = pd.to_datetime(teks, errors="coerce", format=fmt)
        n = int(t.notna().sum())
        if n > jml_terbaik:
            terbaik, jml_terbaik = t, n
        if n == len(teks):
            return t
    # Tidak ada format baku yang cocok sepenuhnya — serahkan ke penebak pandas.
    bebas = pd.to_datetime(teks, errors="coerce")
    return bebas if int(bebas.notna().sum()) > jml_terbaik else terbaik


def cari_kolom_tanggal(df: pd.DataFrame):
    """Cari kolom tanggal kunjungan — bukan tanggal lahir."""
    for kandidat in ("Tgl", "Tanggal", "tanggal_kunjungan", "Tgl Kunjungan", "Tanggal Kunjungan"):
        if kandidat in df.columns:
            return kandidat
    for c in df.columns:
        nama = str(c)
        if re.search(r"tgl|tanggal|date", nama, re.I) and not re.search(r"lahir|birth|daftar", nama, re.I):
            return c
    return None


@st.cache_data(show_spinner=False)
def gabung_berkas(berkas: list, buang_kembar: bool = True):
    """Gabungkan beberapa berkas menjadi satu, urut menurut tanggal kunjungan.

    `berkas` berupa daftar pasangan (nama, isi byte) supaya hasilnya bisa
    di-cache — objek unggahan Streamlit sendiri tidak bisa di-hash.
    Mengembalikan (df_gabungan, ringkasan_per_berkas, jumlah_kembar, nama_kolom_tanggal).
    """
    potongan, ringkasan = [], []

    for nama, data in berkas:
        df = baca_berkas(nama, data)
        df.columns = [str(c).strip() for c in df.columns]
        kol_tgl = cari_kolom_tanggal(df)
        tgl = (_ke_tanggal(df[kol_tgl]) if kol_tgl
               else pd.Series(pd.NaT, index=df.index))
        ringkasan.append({
            "Berkas": nama,
            "Baris": len(df),
            "Kolom": len(df.columns),
            "Tanggal awal": tgl.min().strftime("%d/%m/%Y") if tgl.notna().any() else "—",
            "Tanggal akhir": tgl.max().strftime("%d/%m/%Y") if tgl.notna().any() else "—",
            "Tanggal tak terbaca": int(tgl.isna().sum()),
        })
        df = df.copy()
        df["_urut_tanggal"] = tgl
        df["_asal_berkas"] = nama
        potongan.append(df)

    gab = pd.concat(potongan, ignore_index=True, sort=False)

    kolom_asli = [c for c in gab.columns if not c.startswith("_")]
    kembar = int(gab.duplicated(subset=kolom_asli).sum())
    if buang_kembar and kembar:
        gab = gab.drop_duplicates(subset=kolom_asli).reset_index(drop=True)

    # Urutkan menurut tanggal; baris tanpa tanggal ditaruh di akhir agar terlihat.
    gab = gab.sort_values("_urut_tanggal", kind="stable", na_position="last").reset_index(drop=True)

    kol_tgl_akhir = cari_kolom_tanggal(gab.drop(columns=["_urut_tanggal", "_asal_berkas"]))
    return gab, ringkasan, kembar, kol_tgl_akhir


def ke_format_sigap(df: pd.DataFrame):
    """Ubah data gabungan menjadi 8 kolom masukan SIGAP-Bojonegoro.

    Kolom identitas — nama, NIK, alamat, nomor penjamin — sengaja tidak ikut.
    Mengembalikan (df_sigap, daftar_kolom_yang_tidak_ditemukan).
    """
    ada = {asal: baru for asal, baru in KOLOM_SIGAP.items() if asal in df.columns}
    kurang = [asal for asal in KOLOM_SIGAP if asal not in df.columns]
    if not ada:
        return None, kurang

    out = df[list(ada.keys())].rename(columns=ada).copy()

    if "tanggal_kunjungan" in out.columns:
        out["tanggal_kunjungan"] = _ke_tanggal(out["tanggal_kunjungan"]).dt.strftime("%Y-%m-%d")
    if "umur" in out.columns:
        # "69 Thn 11 Bln 28 Hari" -> 69
        out["umur"] = out["umur"].astype(str).str.extract(r"(\d+)")[0]
    if "no_rm" in out.columns:
        out["no_rm"] = (out["no_rm"].astype(str)
                        .str.replace(r"[\s,]+$", "", regex=True).str.strip())

    urutan = [baru for baru in KOLOM_SIGAP.values() if baru in out.columns]
    return out[urutan], kurang


def ke_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


# ─────────────────────────────────────────────
#  SIDEBAR
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("""
    <div style="padding: 1rem 0 1.5rem;">
        <div style="font-size:1.35rem; font-weight:800; color:#f0f9ff; letter-spacing:-0.03em;">🧹 Data Cleaning SIGAP</div>
        <div style="font-size:0.75rem; color:#475569; margin-top:0.2rem;">Penyiap data untuk SIGAP-Bojonegoro · v2.0</div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("**📂 Unggah Berkas**")
    uploaded_files = st.file_uploader(
        "Pilih satu atau beberapa berkas bulanan",
        type=["xlsx", "xls", "csv"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )
    st.caption("Beberapa berkas bulanan bisa diunggah sekaligus; hasilnya digabung dan diurutkan menurut tanggal kunjungan.")

    st.markdown("---")
    st.markdown("**🔗 Penggabungan**")
    opt_buang_kembar = st.checkbox("Buang baris kembar antarberkas", value=True)
    opt_urut_tanggal = st.checkbox("Urutkan menurut tanggal kunjungan", value=True)
    opt_kolom_asal   = st.checkbox("Tambah kolom Asal Berkas", value=False)

    st.markdown("---")
    st.markdown("**⚙️ Opsi Pembersihan**")

    opt_remove_dup   = st.checkbox("Hapus baris duplikat", value=True)
    opt_strip_comma  = st.checkbox("Hapus trailing koma/spasi", value=True)
    opt_strip_spaces = st.checkbox("Hapus spasi dalam NIK & No Penjamin", value=True)
    opt_format_dates = st.checkbox("Format tanggal ke dd/mm/yyyy", value=True)
    opt_std_ket      = st.checkbox("Standarisasi kolom Keterangan", value=True)
    opt_upper_ket    = st.checkbox("↳ UPPERCASE Keterangan", value=True)
    opt_upper_nama   = st.checkbox("UPPERCASE kolom Nama", value=True)
    opt_title_desa   = st.checkbox("Title Case kolom Desa", value=True)

    st.markdown("---")
    st.markdown("**🔄 Nilai Pengganti**")
    fill_keterangan = st.text_input("Keterangan kosong/simbol", value="Tidak Ada Keterangan")
    fill_pekerjaan  = st.text_input("Pekerjaan kosong", value="Tidak Diketahui")
    fill_rm_lama    = st.text_input("RM Lama kosong", value="-")
    fill_desa       = st.text_input("Desa kosong", value="-")
    fill_no_penjamin = st.text_input("No Penjamin kosong", value="-")

    st.markdown("---")
    st.markdown("**📤 Keluaran**")
    st.caption("Setiap hasil tersedia dalam dua bentuk: arsip lengkap dan berkas siap unggah ke SIGAP, masing-masing sebagai .xlsx maupun .csv.")
    sheet_name = st.text_input("Nama sheet (Excel)", value="Data Bersih")


# ─────────────────────────────────────────────
#  MAIN CONTENT
# ─────────────────────────────────────────────
st.markdown("""
<div class="main-header">
    <h1>🧹 Data Cleaning SIGAP Bojonegoro</h1>
    <p>Membersihkan dan menggabungkan berkas RME bulanan menjadi satu berkas siap pakai</p>
</div>
""", unsafe_allow_html=True)

if not uploaded_files:
    st.markdown("""
    <div class="upload-hint">
        <div style="font-size:2.5rem; margin-bottom:0.75rem;">📂</div>
        <div style="color:#7dd3fc; font-weight:600; font-size:1rem; margin-bottom:0.4rem;">
            Unggah berkas RME di sidebar kiri — boleh beberapa bulan sekaligus
        </div>
        <div style="font-size:0.8rem; color:#475569;">
            Format yang didukung: .xlsx · .xls · .csv<br>
            Termasuk berkas .xls dari RME yang sebenarnya berisi teks berpemisah tab
        </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="section-card">
        <div class="section-title">📋 Fitur Pembersihan</div>
    """, unsafe_allow_html=True)

    cols = st.columns(3)
    features = [
        ("🔗", "Gabung Banyak Berkas", "Beberapa berkas bulanan menjadi satu, urut tanggal"),
        ("📤", "Dua Bentuk Unduhan", "Arsip lengkap, atau 8 kolom siap unggah ke SIGAP"),
        ("🗑️", "Hapus Duplikat", "Deteksi dan hapus baris yang persis sama"),
        ("✂️", "Trailing Koma", "Hapus karakter ' ,' di akhir sel (No RM, NIK, dll)"),
        ("📅", "Format Tanggal", "Ubah ke format dd/mm/yyyy secara konsisten"),
        ("📝", "Standarisasi Keterangan", "Ubah '-' dan '=' menjadi teks bermakna"),
        ("🔠", "Normalisasi Teks", "UPPERCASE nama, Title Case desa"),
        ("📋", "Isi Nilai Kosong", "Isi kolom kosong dengan nilai default"),
    ]
    for i, (icon, title, desc) in enumerate(features):
        with cols[i % 3]:
            st.markdown(f"""
            <div style="background:rgba(30,41,59,0.6); border:1px solid rgba(99,179,237,0.15);
                        border-radius:10px; padding:1rem; margin-bottom:0.75rem; text-align:center;">
                <div style="font-size:1.8rem;">{icon}</div>
                <div style="color:#e2e8f0; font-weight:700; font-size:0.9rem; margin:0.4rem 0 0.3rem;">{title}</div>
                <div style="color:#64748b; font-size:0.78rem;">{desc}</div>
            </div>
            """, unsafe_allow_html=True)

    st.markdown("</div>", unsafe_allow_html=True)
    st.stop()


# ─────────────────────────────────────────────
#  LOAD DATA
# ─────────────────────────────────────────────
berkas_masuk = [(f.name, f.getvalue()) for f in uploaded_files]

try:
    with st.spinner(f"Membaca dan menggabungkan {len(berkas_masuk)} berkas..."):
        df_raw, ringkasan, jml_kembar, kol_tgl = gabung_berkas(berkas_masuk, opt_buang_kembar)
except Exception as e:
    st.error(f"❌ Gagal membaca berkas: {e}")
    st.stop()

asal_berkas = df_raw["_asal_berkas"].copy()
urut_tanggal = df_raw["_urut_tanggal"].copy()
df_raw = df_raw.drop(columns=["_urut_tanggal", "_asal_berkas"])
if opt_kolom_asal:
    df_raw["Asal Berkas"] = asal_berkas.values

# ── Ringkasan penggabungan ──────────────────────────────────
if len(berkas_masuk) > 1 or jml_kembar:
    with st.expander(f"🔗 Ringkasan penggabungan — {len(berkas_masuk)} berkas, {len(df_raw):,} baris".replace(",", "."), expanded=True):
        st.dataframe(pd.DataFrame(ringkasan), use_container_width=True, hide_index=True)
        pesan = []
        if kol_tgl:
            pesan.append(f"Diurutkan menurut kolom **{kol_tgl}**.")
        else:
            pesan.append("⚠️ Kolom tanggal kunjungan tidak ditemukan, urutan mengikuti urutan unggah.")
        if jml_kembar:
            pesan.append(
                f"**{jml_kembar:,}** baris kembar ditemukan dan **dibuang**.".replace(",", ".")
                if opt_buang_kembar else
                f"**{jml_kembar:,}** baris kembar ditemukan dan **dipertahankan**.".replace(",", ".")
            )
        else:
            pesan.append("Tidak ada baris kembar antarberkas.")
        tak_terbaca = int(urut_tanggal.isna().sum())
        if tak_terbaca:
            pesan.append(f"⚠️ **{tak_terbaca}** baris tanggalnya tidak terbaca dan ditaruh di urutan paling akhir.")
        st.markdown(" ".join(pesan))


# ─────────────────────────────────────────────
#  DETECT ISSUES
# ─────────────────────────────────────────────
issues = detect_issues(df_raw)
total_nulls = sum(issues.get("nulls", {}).values())
dup_count   = issues.get("duplicates", 0)

# Metric cards
st.markdown(f"""
<div class="metric-row">
    <div class="metric-card">
        <div class="label">Total Baris</div>
        <div class="value">{len(df_raw):,}</div>
    </div>
    <div class="metric-card">
        <div class="label">Total Kolom</div>
        <div class="value">{len(df_raw.columns)}</div>
    </div>
    <div class="metric-card">
        <div class="label">Duplikat</div>
        <div class="value {'bad' if dup_count else 'good'}">{dup_count}</div>
    </div>
    <div class="metric-card">
        <div class="label">Nilai Kosong</div>
        <div class="value {'bad' if total_nulls else 'good'}">{total_nulls:,}</div>
    </div>
</div>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  TWO COLUMNS: Issues + Preview
# ─────────────────────────────────────────────
col_left, col_right = st.columns([1, 2])

with col_left:
    st.markdown('<div class="section-card">', unsafe_allow_html=True)
    st.markdown('<div class="section-title">🔍 Masalah Terdeteksi</div>', unsafe_allow_html=True)

    if not issues:
        st.markdown('<div class="log-item"><span class="log-icon">✅</span> Data sudah bersih!</div>', unsafe_allow_html=True)
    else:
        if "duplicates" in issues:
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">🔴</span>
                <span><b style="color:#f87171">{issues['duplicates']} baris duplikat</b></span>
            </div>""", unsafe_allow_html=True)

        if "trailing_comma" in issues:
            cols_str = ", ".join(issues["trailing_comma"][:4])
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">🟡</span>
                <span>Trailing koma:<br><span class="badge badge-yellow">{cols_str}</span></span>
            </div>""", unsafe_allow_html=True)

        if "nulls" in issues:
            null_items = "".join([
                f'<span class="badge badge-red" style="margin:2px;">{c}: {v}</span> '
                for c, v in list(issues["nulls"].items())[:8]
            ])
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">🔴</span>
                <span>Nilai kosong:<br>{null_items}</span>
            </div>""", unsafe_allow_html=True)

        if "placeholders" in issues:
            ph_str = ", ".join(issues["placeholders"][:4])
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">🟡</span>
                <span>Nilai placeholder (-/=):<br>
                <span class="badge badge-yellow">{ph_str}</span></span>
            </div>""", unsafe_allow_html=True)

        if "date_cols" in issues:
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">🔵</span>
                <span>Kolom tanggal: <span class="badge badge-blue">{', '.join(issues['date_cols'])}</span></span>
            </div>""", unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)

    # Column selector for dropping
    st.markdown('<div class="section-card">', unsafe_allow_html=True)
    st.markdown('<div class="section-title">🗂️ Pilih Kolom</div>', unsafe_allow_html=True)
    cols_to_drop = st.multiselect(
        "Kolom yang ingin dihapus",
        options=list(df_raw.columns),
        default=[],
        placeholder="Pilih kolom..."
    )
    sort_col = st.selectbox(
        "Urutkan berdasarkan",
        options=["(tidak diurutkan)"] + list(df_raw.columns),
        index=0
    )
    st.markdown('</div>', unsafe_allow_html=True)

with col_right:
    st.markdown('<div class="section-card">', unsafe_allow_html=True)
    st.markdown('<div class="section-title">👁️ Preview Data Asli (10 baris pertama)</div>', unsafe_allow_html=True)
    st.dataframe(df_raw.head(10), use_container_width=True, height=280)
    st.markdown('</div>', unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  CLEAN BUTTON
# ─────────────────────────────────────────────
st.markdown("---")
btn_col, _ = st.columns([1, 3])
with btn_col:
    run_clean = st.button("🚀 Bersihkan Data Sekarang", use_container_width=True)

if run_clean or "df_clean" in st.session_state:

    if run_clean:
        # Build options dict
        strip_space_cols = []
        for col in df_raw.select_dtypes(include="object").columns:
            if "nik" in col.lower() or "penjamin" in col.lower():
                strip_space_cols.append(col)

        fill_nulls_map = {}
        if fill_pekerjaan and "Pekerjaan" in df_raw.columns:
            fill_nulls_map["Pekerjaan"] = fill_pekerjaan
        if fill_rm_lama and "RM Lama" in df_raw.columns:
            fill_nulls_map["RM Lama"] = fill_rm_lama
        if fill_desa and "Desa" in df_raw.columns:
            fill_nulls_map["Desa"] = fill_desa
        if fill_no_penjamin and "No Penjamin" in df_raw.columns:
            fill_nulls_map["No Penjamin"] = fill_no_penjamin

        rename_map = {}
        if "kategori" in df_raw.columns:
            rename_map["kategori"] = "Kategori"

        options = {
            "remove_duplicates":     opt_remove_dup,
            "strip_trailing_comma":  opt_strip_comma,
            "strip_spaces":          opt_strip_spaces,
            "strip_spaces_cols":     strip_space_cols,
            "format_dates":          opt_format_dates,
            "date_format":           "%d/%m/%Y",
            "standardize_keterangan":opt_std_ket,
            "placeholder_replacement": fill_keterangan,
            "uppercase_keterangan":  opt_upper_ket,
            "uppercase_nama":        opt_upper_nama,
            "titlecase_desa":        opt_title_desa,
            "fill_nulls":            fill_nulls_map,
            "rename_cols":           rename_map,
            "drop_cols":             cols_to_drop,
            "sort_by":               sort_col if sort_col != "(tidak diurutkan)" else None,
        }

        with st.spinner("Membersihkan data..."):
            df_clean, clean_log = clean_dataframe(df_raw, options)
        st.session_state["df_clean"] = df_clean
        st.session_state["clean_log"] = clean_log

    df_clean = st.session_state["df_clean"]
    clean_log = st.session_state.get("clean_log", [])

    # ── Results ────────────────────────────────
    removed_rows = len(df_raw) - len(df_clean)
    st.markdown(f"""
    <div class="metric-row">
        <div class="metric-card">
            <div class="label">Baris Tersisa</div>
            <div class="value good">{len(df_clean):,}</div>
        </div>
        <div class="metric-card">
            <div class="label">Baris Dihapus</div>
            <div class="value {'bad' if removed_rows else 'good'}">{removed_rows}</div>
        </div>
        <div class="metric-card">
            <div class="label">Kolom Tersisa</div>
            <div class="value">{len(df_clean.columns)}</div>
        </div>
        <div class="metric-card">
            <div class="label">Nilai Kosong</div>
            <div class="value good">{int(df_clean.isnull().sum().sum())}</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    res_left, res_right = st.columns([1, 2])

    with res_left:
        st.markdown('<div class="section-card">', unsafe_allow_html=True)
        st.markdown('<div class="section-title">✅ Log Pembersihan</div>', unsafe_allow_html=True)
        for icon, msg in clean_log:
            st.markdown(f"""
            <div class="log-item">
                <span class="log-icon">{icon}</span>
                <span>{msg}</span>
            </div>""", unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    with res_right:
        st.markdown('<div class="section-card">', unsafe_allow_html=True)
        st.markdown('<div class="section-title">✨ Preview Data Bersih</div>', unsafe_allow_html=True)
        st.dataframe(df_clean.head(15), use_container_width=True, height=320)
        st.markdown('</div>', unsafe_allow_html=True)

    # ── Unduhan ────────────────────────────────
    st.markdown("---")

    if opt_urut_tanggal:
        kol_tgl_bersih = cari_kolom_tanggal(df_clean)
        if kol_tgl_bersih:
            _t = _ke_tanggal(df_clean[kol_tgl_bersih])
            df_clean = (df_clean.assign(_t=_t)
                        .sort_values("_t", kind="stable", na_position="last")
                        .drop(columns="_t").reset_index(drop=True))

    # Nama berkas mengikuti rentang tanggal isinya, bukan nama berkas asal.
    _t = (_ke_tanggal(df_clean[cari_kolom_tanggal(df_clean)])
          if cari_kolom_tanggal(df_clean) else pd.Series([], dtype="datetime64[ns]"))
    if len(_t) and _t.notna().any():
        base_name = f"RME_{_t.min():%Y-%m-%d}_sd_{_t.max():%Y-%m-%d}"
    else:
        base_name = "RME_gabungan"

    df_sigap, kolom_kurang = ke_format_sigap(df_clean)

    st.markdown('<div class="section-title">📤 Unduh Hasil</div>', unsafe_allow_html=True)
    u1, u2 = st.columns(2)

    with u1:
        st.markdown(f"**📚 Arsip lengkap** — {len(df_clean.columns)} kolom, {len(df_clean):,} baris".replace(",", "."))
        st.caption("Seluruh kolom apa adanya, termasuk identitas pasien. Simpan di tempat yang aman.")
        st.download_button(
            "⬇️ Arsip lengkap (.xlsx)",
            data=to_excel_bytes(df_clean, sheet_name),
            file_name=f"{base_name}_lengkap.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        st.download_button(
            "⬇️ Arsip lengkap (.csv)",
            data=ke_csv_bytes(df_clean.fillna("-")),
            file_name=f"{base_name}_lengkap.csv",
            mime="text/csv",
            use_container_width=True,
        )

    with u2:
        if df_sigap is None:
            st.warning("Kolom yang dibutuhkan SIGAP tidak ditemukan pada data ini.")
        else:
            st.markdown(f"**🏥 Siap unggah ke SIGAP** — {len(df_sigap.columns)} kolom, {len(df_sigap):,} baris".replace(",", "."))
            st.caption("Tanpa nama, NIK, alamat, dan nomor penjamin — aman dibawa ke luar puskesmas.")
            st.download_button(
                "⬇️ Siap SIGAP (.csv)",
                data=ke_csv_bytes(df_sigap),
                file_name=f"{base_name}_sigap.csv",
                mime="text/csv",
                use_container_width=True,
                type="primary",
            )
            st.download_button(
                "⬇️ Siap SIGAP (.xlsx)",
                data=to_excel_bytes(df_sigap, "Data SIGAP"),
                file_name=f"{base_name}_sigap.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
            if kolom_kurang:
                st.caption("Kolom asal yang tidak ditemukan: " + ", ".join(kolom_kurang))
            with st.expander("Pratinjau data siap SIGAP"):
                st.dataframe(df_sigap.head(15), use_container_width=True, hide_index=True)

    # Null checker after clean
    remaining_nulls = df_clean.isnull().sum()
    remaining_nulls = remaining_nulls[remaining_nulls > 0]
    if not remaining_nulls.empty:
        with st.expander("⚠️ Nilai kosong yang masih tersisa"):
            st.dataframe(
                remaining_nulls.reset_index().rename(columns={"index": "Kolom", 0: "Jumlah Kosong"}),
                use_container_width=True
            )
